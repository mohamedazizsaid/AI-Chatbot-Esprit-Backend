import PyPDF2
from PIL import Image
import pytesseract
from docx import Document
import openpyxl
from typing import List, Dict
import os

class DocumentProcessor:
    def __init__(self, chunk_size: int = 500, chunk_overlap: int = 50):
        """
        Initialise le processeur de documents
        chunk_size: Taille des chunks de texte
        chunk_overlap: Chevauchement entre chunks
        """
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
        
    def process_document(self, file_path: str, filename: str) -> List[Dict]:
        """Traite un document et retourne des chunks"""
        
        file_extension = os.path.splitext(filename)[1].lower()
        
        # Extraction du texte selon le type de fichier
        if file_extension == '.pdf':
            text = self._extract_from_pdf(file_path)
        elif file_extension in ['.png', '.jpg', '.jpeg', '.bmp', '.tiff']:
            text = self._extract_from_image(file_path)
        elif file_extension == '.docx':
            text = self._extract_from_docx(file_path)
        elif file_extension == '.txt':
            text = self._extract_from_txt(file_path)
        elif file_extension in ['.xlsx', '.xls']:
            text = self._extract_from_excel(file_path)
        else:
            raise ValueError(f"Type de fichier non supporté: {file_extension}")
        
        # Découper en chunks
        chunks = self._create_chunks(text, filename)
        
        return chunks
    
    def _extract_from_pdf(self, file_path: str) -> str:
        """Extrait le texte d'un PDF"""
        text = ""
        try:
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                for page_num, page in enumerate(pdf_reader.pages):
                    page_text = page.extract_text()
                    text += f"\n[Page {page_num + 1}]\n{page_text}"
        except Exception as e:
            print(f"Erreur lors de l'extraction du PDF: {e}")
        
        return text
    
    def _extract_from_image(self, file_path: str) -> str:
        """Extrait le texte d'une image via OCR"""
        try:
            image = Image.open(file_path)
            # Configuration OCR pour le français
            text = pytesseract.image_to_string(image, lang='fra')
            return text
        except Exception as e:
            print(f"Erreur OCR: {e}")
            return ""
    
    def _extract_from_docx(self, file_path: str) -> str:
        """Extrait le texte d'un document Word"""
        try:
            doc = Document(file_path)
            text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
            return text
        except Exception as e:
            print(f"Erreur lors de l'extraction du DOCX: {e}")
            return ""
    
    def _extract_from_txt(self, file_path: str) -> str:
        """Extrait le texte d'un fichier texte"""
        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                return file.read()
        except Exception as e:
            print(f"Erreur lors de l'extraction du TXT: {e}")
            return ""
    
    def _extract_from_excel(self, file_path: str) -> str:
        """Extrait le texte d'un fichier Excel"""
        try:
            workbook = openpyxl.load_workbook(file_path)
            text = ""
            
            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                text += f"\n[Feuille: {sheet_name}]\n"
                
                for row in sheet.iter_rows(values_only=True):
                    row_text = " | ".join([str(cell) if cell is not None else "" for cell in row])
                    text += row_text + "\n"
            
            return text
        except Exception as e:
            print(f"Erreur lors de l'extraction d'Excel: {e}")
            return ""
    
    def _create_chunks(self, text: str, filename: str) -> List[Dict]:
        """Découpe le texte en chunks avec metadata"""
        chunks = []
        
        # Nettoyer le texte
        text = text.strip()
        
        if not text:
            return chunks
        
        # Découpage simple par caractères
        start = 0
        chunk_id = 0
        
        while start < len(text):
            end = start + self.chunk_size
            
            # Essayer de couper à un espace ou point
            if end < len(text):
                # Chercher le dernier espace ou point
                last_space = text.rfind(' ', start, end)
                last_period = text.rfind('.', start, end)
                
                cutoff = max(last_space, last_period)
                if cutoff > start:
                    end = cutoff + 1
            
            chunk_text = text[start:end].strip()
            
            if chunk_text:
                chunks.append({
                    "text": chunk_text,
                    "metadata": {
                        "filename": filename,
                        "chunk_id": chunk_id,
                        "start_char": start,
                        "end_char": end
                    }
                })
            
            chunk_id += 1
            start = end - self.chunk_overlap
        
        return chunks
    
    def get_supported_formats(self) -> List[str]:
        """Retourne la liste des formats supportés"""
        return ['.pdf', '.docx', '.txt', '.xlsx', '.xls', '.png', '.jpg', '.jpeg', '.bmp', '.tiff']